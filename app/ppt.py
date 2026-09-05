"""知识卡片导出：把文章卡片渲染成 16:9 PPTX。

版式尽量复用前端 app/static/index.html 里的知识卡片结构：
第 1 页 = 文献期刊信息 + 文章主题 + 典型反应式（仅图）+ 底物范围 + 关键影响因素 + 分类；
第 2 页 = 反应机理（原文章有图则取图，附带机理文字）。
"""

from __future__ import annotations

from io import BytesIO
from typing import Any, Dict, List, Optional

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from . import config
from .db import get_article, list_articles

# ---- 配色（与前端一致） ----
BLUE = RGBColor(0x1E, 0x6F, 0xD0)
INK = RGBColor(0x1F, 0x27, 0x33)
MUTED = RGBColor(0x6B, 0x7A, 0x89)
LINE = RGBColor(0xE4, 0xE8, 0xEE)
BOX_BG = RGBColor(0xFB, 0xFC, 0xFE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SW, SH = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.55)


def _num(article_id: int) -> str:
    return "KC-" + str(article_id).zfill(4)


def _card_no(article: Dict[str, Any]) -> str:
    return article.get("card_no") or _num(article.get("id", 0))


def _meta(art: Dict[str, Any]) -> str:
    parts = [
        art.get("journal") or "",
        art.get("year") or "",
        ("卷 " + art["volume"]) if art.get("volume") else "",
        ("期 " + art["issue"]) if art.get("issue") else "",
        art.get("pages") or "",
    ]
    return " · ".join(p for p in parts if p)


def _topic(ex: Dict[str, Any]) -> str:
    kw = ex.get("keywords") or {}
    rxs = ex.get("reactions") or []
    rt = (kw.get("reactions") or [None])[0] or (rxs[0].get("name") if rxs else "") or ""
    core = " · ".join((kw.get("tags") or [])[:2]) or (ex.get("article", {}).get("title") or "")
    topic = (rt + " · " + core) if rt else core
    return topic or (ex.get("article", {}).get("title") or "")


def _category(ex: Dict[str, Any]) -> str:
    return ex.get("category") or "其它"


def _img_abs(name: str) -> Optional[str]:
    """图片名 -> 磁盘绝对路径；不存在则返回 None。"""
    if not name:
        return None
    p = config.IMAGES_DIR / name
    return str(p) if p.exists() else None


def _as_dict(obj: Any) -> Dict[str, Any]:
    """pydantic 模型 → 普通 dict（嵌套模型一并转换）。"""
    if obj is None:
        return {}
    if isinstance(obj, dict):
        return obj
    if hasattr(obj, "model_dump"):
        return obj.model_dump()
    return dict(obj) if isinstance(obj, dict) else {}


def _image_src(article: Dict[str, Any], role: str, fallback_idx: Optional[int] = None) -> Optional[str]:
    """按 role 取一张图片路径（与前端图片分配逻辑一致）。"""
    imgs = article.get("images") or []
    if not imgs:
        return None
    for im in imgs:
        if im.get("role") == role:
            path = _img_abs(im.get("path"))
            if path:
                return path
    if fallback_idx is not None and fallback_idx < len(imgs):
        return _img_abs(imgs[fallback_idx].get("path"))
    return None


def _reaction_images(article: Dict[str, Any]) -> List[str]:
    """只取第一张主反应式图；条件优化/其它图不放进主反应栏。"""
    imgs = article.get("images") or []
    rx = []
    seen = set()
    for im in imgs:
        if im.get("role") == "reaction":
            path = _img_abs(im.get("path"))
            if path and path not in seen:
                rx.append(path)
                seen.add(path)
            if len(rx) >= 1:
                break
    if not rx:
        for im in imgs:
            path = _img_abs(im.get("path"))
            if path:
                rx.append(path)
                break
    return rx


# ---------- 底层绘制 ----------


def _set_text(tf, text: str, size: float, color=MUTED, bold=False, align=PP_ALIGN.LEFT,
              anchor=MSO_ANCHOR.TOP) -> None:
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold


def _add_box(slide, left, top, width, height, fill=WHITE, line=LINE, radius=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    if radius:
        shp.adjustments[0] = radius
    return shp


def _add_text(slide, left, top, width, height, text, size, color=MUTED, bold=False,
              align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, fill=None, line=None, radius=None):
    shp = _add_box(slide, left, top, width, height, fill=fill or WHITE, line=line, radius=radius)
    tf = shp.text_frame
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.08)
    tf.margin_bottom = Inches(0.08)
    _set_text(tf, text, size, color, bold, align, anchor)
    return shp


def _add_label_value(slide, left, top, width, height, label, value, size=9.5):
    box = _add_box(slide, left, top, width, height, fill=WHITE, line=LINE, radius=0.08)
    tf = box.text_frame
    tf.margin_left = Inches(0.16)
    tf.margin_right = Inches(0.16)
    tf.margin_top = Inches(0.10)
    tf.margin_bottom = Inches(0.10)
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = label
    r.font.size = Pt(size)
    r.font.color.rgb = INK
    r.font.bold = True
    if value:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.LEFT
        r2 = p2.add_run()
        r2.text = value
        r2.font.size = Pt(9)
        r2.font.color.rgb = MUTED
    return box


def _fit_picture(slide, img_path: str, left, top, width, height):
    """按原比例缩放图片，居中放进指定方框。"""
    try:
        with Image.open(img_path) as im:
            iw, ih = im.size
    except Exception:
        return None
    if not iw or not ih:
        return None
    scale = min(width / iw, height / ih)
    nw, nh = int(iw * scale), int(ih * scale)
    nleft = left + int((width - nw) / 2)
    ntop = top + int((height - nh) / 2)
    try:
        pic = slide.shapes.add_picture(img_path, nleft, ntop, max(nw, 1), max(nh, 1))
        return pic
    except Exception:
        return None


def _add_fig_placeholder(slide, left, top, width, height, text):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = BOX_BG
    shp.line.color.rgb = RGBColor(0xC9, 0xD2, 0xDC)
    shp.line.dash_style = 4  # dashed
    shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    _set_text(tf, text, 9, RGBColor(0x9A, 0xA7, 0xB5), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return shp


def _add_picture_or_placeholder(slide, img_path, left, top, width, height, placeholder):
    if img_path and _fit_picture(slide, img_path, left, top, width, height):
        return
    _add_fig_placeholder(slide, left, top, width, height, placeholder)


# ---------- 页码 ----------


def _add_header(slide, article: Dict[str, Any], meta: str):
    """顶部三栏：文献期刊信息 / 抓取时间 / 知识卡片序号。"""
    row = Inches(0.55)
    h = Inches(0.85)
    _add_label_value(slide, MARGIN, row, Inches(7.6), h,
                     "文献期刊信息", meta or "（未识别期刊信息）")
    _add_label_value(slide, MARGIN + Inches(7.8), row, Inches(2.1), h,
                     "抓取时间", article.get("created_at") or "")
    _add_label_value(slide, MARGIN + Inches(10.1), row, Inches(2.1), h,
                     "知识卡片序号", _card_no(article))


def _add_topic(slide, ex: Dict[str, Any]):
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                 MARGIN, Inches(1.6), SW - MARGIN * 2, Inches(0.75))
    box.fill.solid()
    box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = LINE
    box.line.width = Pt(0.75)
    box.shadow.inherit = False
    tf = box.text_frame
    tf.margin_left = Inches(0.16)
    tf.margin_top = Inches(0.12)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "文章主题"
    r.font.size = Pt(9.5)
    r.font.color.rgb = MUTED
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = _topic(ex)
    r2.font.size = Pt(13)
    r2.font.color.rgb = INK
    r2.font.bold = True


def _add_footer(slide, ex: Dict[str, Any], page_label: str = ""):
    left = MARGIN
    top = SH - Inches(0.6)
    r = _add_text(slide, left, top, Inches(5), Inches(0.42), "",
                  9, MUTED, align=PP_ALIGN.LEFT)
    tf = r.text_frame
    p = tf.paragraphs[0]
    p.add_run().text = "分类：" + _category(ex)
    r2 = _add_text(slide, left + Inches(6), top, Inches(5.8), Inches(0.42), "",
                   9, MUTED, align=PP_ALIGN.RIGHT)
    p2 = r2.text_frame.paragraphs[0]
    p2.add_run().text = "化学文献库 · " + (page_label or _card_no(ex.get("_article", {})))


# ---------- 幻灯片 ----------


def _slide_main(prs, article: Dict[str, Any]) -> None:
    ex = _as_dict(article.get("extracted"))
    art = ex.get("article") or {}
    ex["_article"] = article
    scope = ex.get("substrate_scope") or {}

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = RGBColor(0xF4, 0xF6, 0xF9)

    _add_header(slide, article, _meta(art))
    _add_topic(slide, ex)

    grid_top = Inches(2.5)
    grid_h = SH - grid_top - Inches(0.7)
    # 左列：典型反应式（占两行）
    col_left_w = Inches(6.6)
    rx_box_left = MARGIN
    rx_box_right = MARGIN + col_left_w
    react_box = _add_box(slide, rx_box_left, grid_top, col_left_w, grid_h, fill=WHITE, line=LINE, radius=0.05)
    tf = react_box.text_frame
    tf.margin_left = Inches(0.16)
    tf.margin_top = Inches(0.14)
    tf.margin_right = Inches(0.16)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "文章的典型反应式（取最关键 1–2 个，仅图）"
    r.font.size = Pt(10)
    r.font.color.rgb = INK
    r.font.bold = True
    rx_imgs = _reaction_images(article)
    fig_top = grid_top + Inches(0.45)
    fig_bottom = grid_top + grid_h - Inches(0.16)
    if len(rx_imgs) == 1:
        _add_picture_or_placeholder(slide, rx_imgs[0],
                                    rx_box_left + Inches(0.2), fig_top,
                                    col_left_w - Inches(0.4), fig_bottom - fig_top,
                                    "未提取到典型反应式图")
    elif len(rx_imgs) >= 1:
        _add_picture_or_placeholder(slide, rx_imgs[0],
                                    rx_box_left + Inches(0.2), fig_top,
                                    col_left_w - Inches(0.4), fig_bottom - fig_top,
                                    "未提取到典型反应式图")
    else:
        _add_fig_placeholder(slide, rx_box_left + Inches(0.2), fig_top,
                             col_left_w - Inches(0.4), fig_bottom - fig_top,
                             "未提取到典型反应式图")

    # 右列上下两盒
    right_left = rx_box_right + Inches(0.18)
    right_w = SW - MARGIN - right_left
    gap = Inches(0.14)
    half_h = (grid_h - gap) / 2
    # 底物适用范围
    scope_box = _add_box(slide, right_left, grid_top, right_w, half_h, fill=WHITE, line=LINE, radius=0.05)
    stf = scope_box.text_frame
    stf.margin_left = Inches(0.16)
    stf.margin_top = Inches(0.14)
    sp = stf.paragraphs[0]
    sr = sp.add_run()
    sr.text = "底物适用范围"
    sr.font.size = Pt(10)
    sr.font.color.rgb = INK
    sr.font.bold = True
    scope_img = _image_src(article, "scope", fallback_idx=2)
    _add_picture_or_placeholder(slide, scope_img,
                                right_left + Inches(0.2), grid_top + Inches(0.42),
                                right_w - Inches(0.4), half_h - Inches(0.9),
                                "底物范围图")
    summary = (scope.get("summary") or "")[:100]
    if summary:
        tb = slide.shapes.add_textbox(right_left + Inches(0.2),
                                      grid_top + half_h - Inches(0.4),
                                      right_w - Inches(0.4), Inches(0.32))
        _set_text(tb.text_frame, summary, 8.5, MUTED)

    # 关键影响因素
    note_top = grid_top + half_h + gap
    note_box = _add_box(slide, right_left, note_top, right_w, half_h, fill=WHITE, line=LINE, radius=0.05)
    ntf = note_box.text_frame
    ntf.margin_left = Inches(0.16)
    ntf.margin_top = Inches(0.14)
    ntf.margin_right = Inches(0.16)
    np_ = ntf.paragraphs[0]
    nr = np_.add_run()
    nr.text = "对反应的关键影响因素"
    nr.font.size = Pt(10)
    nr.font.color.rgb = INK
    nr.font.bold = True
    notes = (scope.get("notes") or "（未提取到关键影响因素）")[:400]
    np2 = ntf.add_paragraph()
    nr2 = np2.add_run()
    nr2.text = notes
    nr2.font.size = Pt(9)
    nr2.font.color.rgb = INK

    _add_footer(slide, ex)


def _slide_mechanism(prs, article: Dict[str, Any]) -> None:
    ex = _as_dict(article.get("extracted"))
    mech = ex.get("mechanism") or {}
    has_mech_img = any(im.get("role") == "mechanism" and _img_abs(im.get("path"))
                       for im in (article.get("images") or []))
    if not (mech.get("overall") or mech.get("steps") or has_mech_img):
        return
    ex["_article"] = article
    art = ex.get("article") or {}

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = RGBColor(0xF4, 0xF6, 0xF9)

    _add_header(slide, article, _meta(art))
    _add_topic(slide, ex)

    top = Inches(2.5)
    bottom = SH - Inches(0.7)
    mimg = _image_src(article, "mechanism", fallback_idx=3)
    _add_picture_or_placeholder(slide, mimg,
                                MARGIN, top, SW - MARGIN * 2, bottom - top - Inches(0.6),
                                "机理图（原文有图则抓取）")
    overall = (mech.get("overall") or "")[:220]
    if overall:
        tb = slide.shapes.add_textbox(MARGIN, bottom - Inches(0.5),
                                      SW - MARGIN * 2, Inches(0.42))
        _set_text(tb.text_frame, overall, 9, INK)

    _add_footer(slide, ex, page_label="机理页")


def build_presentation(articles: List[Dict[str, Any]]) -> BytesIO:
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    for a in articles:
        _slide_main(prs, a)
        _slide_mechanism(prs, a)
    buf = BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


def export_all_ppt() -> BytesIO:
    ids = [a["id"] for a in list_articles()]
    return build_presentation([get_article(i) for i in ids if get_article(i)])


def export_single_ppt(article_id: int) -> BytesIO:
    a = get_article(article_id)
    if not a:
        raise ValueError(f"文章 #{article_id} 不存在")
    return build_presentation([a])
